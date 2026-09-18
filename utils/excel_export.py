import os
import aiosqlite
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import datetime
from database.local_db import DB_PATH
from config import BRANCH_NAMES

async def generate_excel_report(file_path: str = "CRM_Hisobot.xlsx", branch_id: int | None = None) -> str:
    wb = openpyxl.Workbook()
    
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        
        # 1. Ombor qoldig'i
        ws_ombor = wb.active
        ws_ombor.title = "Ombor Qoldig'i"
        
        headers_ombor = [
            "Filial", "Rulon kodi", "Eni (m)", "Rangi", "Boshlang'ich metr", 
            "Qoldiq metr", "Maydon (m²)", "Tovar qiymati ($)", "Holati", "Kirim sanasi"
        ]
        ws_ombor.append(headers_ombor)
        
        if branch_id is not None:
            q_rolls = "SELECT * FROM rolls WHERE branch_id = ? ORDER BY id DESC"
            p_rolls = (branch_id,)
        else:
            q_rolls = "SELECT * FROM rolls ORDER BY branch_id ASC, id DESC"
            p_rolls = ()
            
        async with db.execute(q_rolls, p_rolls) as cursor:
            rolls = await cursor.fetchall()
            for r in rolls:
                b_name = BRANCH_NAMES.get(r["branch_id"], f"Filial-{r['branch_id']}")
                ws_ombor.append([
                    b_name, r["roll_code"], r["width"], r["color"], r["initial_length"],
                    r["current_length"], r["area_m2"], r["total_price"],
                    "Mavjud" if r["status"] == "active" and r["current_length"] > 0 else "Tugagan",
                    r["created_at"]
                ])
                
        # 2. Sotuvlar tarixi
        ws_sales = wb.create_sheet(title="Sotuvlar Tarixi")
        headers_sales = [
            "Filial", "Sana va Vaqt", "Rulon kodi", "Mahsulot", "Sotilgan metr (m)", 
            "Maydoni (m²)", "Narx ($/m²)", "Jami summa ($)"
        ]
        ws_sales.append(headers_sales)
        
        if branch_id is not None:
            q_sales = "SELECT * FROM sales WHERE branch_id = ? ORDER BY id DESC"
            p_sales = (branch_id,)
        else:
            q_sales = "SELECT * FROM sales ORDER BY branch_id ASC, id DESC"
            p_sales = ()
            
        async with db.execute(q_sales, p_sales) as cursor:
            sales = await cursor.fetchall()
            for s in sales:
                b_name = BRANCH_NAMES.get(s["branch_id"], f"Filial-{s['branch_id']}")
                ws_sales.append([
                    b_name, s["created_at"], s["roll_code"], f"{s['width']}x{s['sold_length']}m - {s['color']}",
                    s["sold_length"], s["area_m2"], s["price_per_m2"], s["total_price"]
                ])
                
        # 3. Kassa tarixi
        ws_cash = wb.create_sheet(title="Kassa Amallari")
        headers_cash = [
            "Filial", "Sana va Vaqt", "Operatsiya turi", "Summa ($)", "Izoh", "Kassa qoldig'i ($)"
        ]
        ws_cash.append(headers_cash)
        
        if branch_id is not None:
            q_cash = "SELECT * FROM cashbox WHERE branch_id = ? ORDER BY id DESC"
            p_cash = (branch_id,)
        else:
            q_cash = "SELECT * FROM cashbox ORDER BY branch_id ASC, id DESC"
            p_cash = ()
            
        async with db.execute(q_cash, p_cash) as cursor:
            cash = await cursor.fetchall()
            for c in cash:
                b_name = BRANCH_NAMES.get(c["branch_id"], f"Filial-{c['branch_id']}")
                ws_cash.append([
                    b_name, c["created_at"],
                    "Kirim (Sotuv)" if c["operation_type"] == "INCOME" else "Chiqim (Topshirildi)",
                    c["amount"], c["note"], c["balance_after"]
                ])

    for ws in [ws_ombor, ws_sales, ws_cash]:
        for col_num in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            
        for row in range(2, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=row, column=col)
                cell.border = thin_border
                if isinstance(cell.value, (int, float)):
                    cell.alignment = right_align
                else:
                    cell.alignment = left_align
                    
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(file_path)
    return file_path
